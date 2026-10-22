#!/usr/bin/env python3
"""Export every worksheet in the input Excel workbooks to UTF-8 CSV."""

import argparse
import csv
from pathlib import Path
import re
import sys

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent


def safe_name(name):
    """Keep worksheet names readable while removing unsafe path characters."""
    return re.sub(r'[<>:"/\\|?*\x00-\x1f]', '_', name).strip(' .') or 'sheet'


def convert_workbook(source, output_dir):
    """Write one CSV per worksheet, including hidden and empty sheets."""
    destination = output_dir / source.stem
    destination.mkdir(parents=True, exist_ok=True)
    workbook = load_workbook(source, read_only=True, data_only=True)
    used_names = set()
    count = 0
    try:
        for sheet in workbook.worksheets:
            name = safe_name(sheet.title)
            candidate = name
            suffix = 2
            while candidate.casefold() in used_names:
                candidate = f'{name}_{suffix}'
                suffix += 1
            used_names.add(candidate.casefold())
            target = destination / f'{candidate}.csv'
            # Ignore headers: Excel adds Column1, Column2, etc. to empty columns.
            # Scan all data rows so values appearing only in later rows survive.
            # Two passes keep large workbooks out of memory.
            populated_columns = set()
            for row in sheet.iter_rows(min_row=2, values_only=True):
                for index, value in enumerate(row):
                    if value is not None and (not isinstance(value, str) or value.strip()):
                        populated_columns.add(index)
            columns = sorted(populated_columns)
            rows = 0
            with target.open('w', encoding='utf-8', newline='') as handle:
                writer = csv.writer(handle)
                for row in sheet.iter_rows(values_only=True):
                    writer.writerow(row[index] for index in columns)
                    rows += 1
            print(f'{source.name} / {sheet.title} -> {target} '
                  f'({rows:,} rows, {len(columns)} columns; '
                  f'dropped {sheet.max_column - len(columns)} empty columns)')
            count += 1
    finally:
        workbook.close()
    return count


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--input-dir', type=Path, default=ROOT / 'input-q3-bulk-data',
                        help='Directory containing .xlsx files (default: input-q3-bulk-data)')
    parser.add_argument('--output-dir', type=Path, default=ROOT / 'output-q3-bulk-data',
                        help='CSV destination (default: output-q3-bulk-data)')
    args = parser.parse_args()
    if not args.input_dir.is_dir():
        parser.error(f'Input directory does not exist: {args.input_dir}')
    sources = sorted(path for path in args.input_dir.iterdir()
                     if path.is_file() and path.suffix.lower() == '.xlsx'
                     and not path.name.startswith('~$'))
    if not sources:
        parser.error(f'No .xlsx files found in {args.input_dir}')
    count = 0
    failures = 0
    for source in sources:
        try:
            count += convert_workbook(source, args.output_dir)
        except Exception as exc:
            print(f'Error converting {source.name}: {exc}', file=sys.stderr)
            failures += 1
    print(f'Wrote {count} CSV files from {len(sources) - failures} workbooks.')
    return 1 if failures else 0


if __name__ == '__main__':
    sys.exit(main())
